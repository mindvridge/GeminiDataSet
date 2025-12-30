#!/usr/bin/env python3
"""
JSONL → Excel 변환 스크립트 (기획자 검토용)

사용법:
    python scripts/convert_to_excel.py --input data/raw/track_b/career --output exports/career.xlsx
    python scripts/convert_to_excel.py --input data/raw/track_b/career/short.jsonl --output exports/career_short.xlsx
"""

import argparse
import json
from pathlib import Path
from datetime import datetime

try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError:
    print("필요한 패키지를 설치합니다...")
    import subprocess
    subprocess.run(["pip", "install", "pandas", "openpyxl"], check=True)
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows


def load_jsonl_files(input_path: Path) -> list:
    """JSONL 파일 로드"""
    sessions = []

    if input_path.is_file():
        files = [input_path]
    else:
        files = sorted(input_path.glob("*.jsonl"))

    for file in files:
        # 파일명에서 타입 추출
        if "short" in file.name:
            turn_type = "short"
        elif "medium" in file.name:
            turn_type = "medium"
        elif "long" in file.name:
            turn_type = "long"
        else:
            turn_type = "unknown"

        with open(file, 'r', encoding='utf-8') as f:
            for line in f:
                if line.strip():
                    try:
                        data = json.loads(line)
                        data['_turn_type'] = turn_type
                        data['_source_file'] = file.name
                        sessions.append(data)
                    except json.JSONDecodeError:
                        continue

    return sessions


def format_conversation(turns: list, include_reasoning: bool = False) -> str:
    """대화 내용을 읽기 쉬운 텍스트로 변환"""
    lines = []
    for i, turn in enumerate(turns, 1):
        client = turn.get('client', {})
        therapist = turn.get('therapist', {})

        # 내담자 발화
        client_text = client.get('text', '')
        emotion = client.get('emotional_state', '')
        lines.append(f"[턴 {i}] 내담자 ({emotion})")
        lines.append(f"{client_text}")
        lines.append("")

        # 상담사 발화
        if include_reasoning:
            reasoning = therapist.get('clinical_reasoning', '')
            strategy = therapist.get('intervention_strategy', '')
            lines.append(f"[상담사 판단] {reasoning}")
            lines.append(f"[개입 전략] {strategy}")

        utterance = therapist.get('utterance', '')
        technique = therapist.get('empathy_technique', '')
        lines.append(f"[턴 {i}] 상담사 ({technique})")
        lines.append(f"{utterance}")
        lines.append("")
        lines.append("-" * 40)
        lines.append("")

    return "\n".join(lines)


def create_summary_sheet(wb: Workbook, sessions: list):
    """요약 시트 생성"""
    ws = wb.active
    ws.title = "요약"

    # 헤더 스타일
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    # 통계 계산
    total = len(sessions)
    by_type = {}
    for s in sessions:
        t = s.get('_turn_type', 'unknown')
        by_type[t] = by_type.get(t, 0) + 1

    # 요약 정보 작성
    ws['A1'] = "상담 데이터 요약 보고서"
    ws['A1'].font = Font(bold=True, size=16)
    ws.merge_cells('A1:D1')

    ws['A3'] = "생성 일시:"
    ws['B3'] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    ws['A4'] = "총 세션 수:"
    ws['B4'] = total

    row = 6
    ws[f'A{row}'] = "타입별 분포"
    ws[f'A{row}'].font = Font(bold=True)

    for turn_type, count in sorted(by_type.items()):
        row += 1
        ws[f'A{row}'] = f"  - {turn_type}"
        ws[f'B{row}'] = f"{count}건"

    # 열 너비 조정
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25


def create_list_sheet(wb: Workbook, sessions: list):
    """세션 목록 시트 생성"""
    ws = wb.create_sheet("세션 목록")

    # 헤더
    headers = ["No", "타입", "Session ID", "내담자 프로필", "주호소", "턴 수", "세션 요약"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # 데이터 추가
    for idx, session in enumerate(sessions, 1):
        ws.cell(row=idx+1, column=1, value=idx)
        ws.cell(row=idx+1, column=2, value=session.get('_turn_type', ''))
        ws.cell(row=idx+1, column=3, value=session.get('session_id', ''))
        ws.cell(row=idx+1, column=4, value=session.get('client_profile', '')[:100])
        ws.cell(row=idx+1, column=5, value=session.get('presenting_problem', '')[:100])
        ws.cell(row=idx+1, column=6, value=len(session.get('turns', [])))
        ws.cell(row=idx+1, column=7, value=session.get('session_summary', '')[:200])

    # 열 너비 조정
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 40
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 8
    ws.column_dimensions['G'].width = 60

    # 필터 추가
    ws.auto_filter.ref = f"A1:G{len(sessions)+1}"


def create_conversation_sheet(wb: Workbook, sessions: list, max_sessions: int = 100):
    """대화 내용 시트 생성 (샘플)"""
    ws = wb.create_sheet("대화 내용 (샘플)")

    ws['A1'] = f"대화 내용 샘플 (처음 {max_sessions}건)"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:C1')

    row = 3
    for idx, session in enumerate(sessions[:max_sessions], 1):
        # 세션 헤더
        ws[f'A{row}'] = f"━━━ 세션 {idx}: {session.get('session_id', '')} ━━━"
        ws[f'A{row}'].font = Font(bold=True, size=12, color="4472C4")
        ws.merge_cells(f'A{row}:C{row}')
        row += 1

        ws[f'A{row}'] = f"타입: {session.get('_turn_type', '')} | 턴 수: {len(session.get('turns', []))}"
        row += 1

        ws[f'A{row}'] = f"주호소: {session.get('presenting_problem', '')}"
        row += 1
        row += 1

        # 대화 내용
        for turn_idx, turn in enumerate(session.get('turns', []), 1):
            client = turn.get('client', {})
            therapist = turn.get('therapist', {})

            # 내담자
            ws[f'A{row}'] = f"[턴 {turn_idx}] 내담자"
            ws[f'A{row}'].font = Font(bold=True, color="E67E22")
            row += 1
            ws[f'A{row}'] = client.get('text', '')
            ws[f'A{row}'].alignment = Alignment(wrap_text=True)
            row += 1

            # 상담사
            ws[f'A{row}'] = f"[턴 {turn_idx}] 상담사"
            ws[f'A{row}'].font = Font(bold=True, color="27AE60")
            row += 1
            ws[f'A{row}'] = therapist.get('utterance', '')
            ws[f'A{row}'].alignment = Alignment(wrap_text=True)
            row += 2

        row += 2

    ws.column_dimensions['A'].width = 100


def convert_to_excel(input_path: str, output_path: str, max_conversation_samples: int = 100):
    """메인 변환 함수"""
    input_path = Path(input_path)
    output_path = Path(output_path)

    print(f"📂 데이터 로드 중: {input_path}")
    sessions = load_jsonl_files(input_path)
    print(f"   총 {len(sessions)}건 로드")

    # 출력 디렉토리 생성
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 워크북 생성
    wb = Workbook()

    print("📊 요약 시트 생성 중...")
    create_summary_sheet(wb, sessions)

    print("📋 세션 목록 시트 생성 중...")
    create_list_sheet(wb, sessions)

    print("💬 대화 내용 시트 생성 중...")
    create_conversation_sheet(wb, sessions, max_conversation_samples)

    # 저장
    wb.save(output_path)
    print(f"\n✅ Excel 파일 저장 완료: {output_path}")
    print(f"   - 요약: 통계 정보")
    print(f"   - 세션 목록: 전체 {len(sessions)}건 (필터 가능)")
    print(f"   - 대화 내용: 샘플 {min(len(sessions), max_conversation_samples)}건")


def main():
    parser = argparse.ArgumentParser(description="JSONL → Excel 변환 (기획자 검토용)")
    parser.add_argument("--input", "-i", required=True, help="입력 JSONL 파일 또는 디렉토리")
    parser.add_argument("--output", "-o", required=True, help="출력 Excel 파일 경로")
    parser.add_argument("--samples", "-s", type=int, default=100, help="대화 내용 샘플 수 (기본값: 100)")

    args = parser.parse_args()
    convert_to_excel(args.input, args.output, args.samples)


if __name__ == "__main__":
    main()
